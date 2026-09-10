from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from app.ingestion.chunking import ChunkerConfig, chunk_document
from app.ingestion.parsers import build_default_registry


def parsed_record(document_factory, content, name):
    parsed = build_default_registry().parse_bytes(content, suffix=Path(name).suffix)
    return parsed, document_factory(
        text=parsed.text,
        sections=parsed.sections,
        tables=parsed.tables,
        parser_name=parsed.parser_name,
        parser_version=parsed.parser_version,
        parse_warnings=parsed.parse_warnings,
        source_path=name,
    )


def word_bytes():
    document = Document()
    document.add_heading("Travel policy", level=1)
    document.add_paragraph("Ordinary employees only. Approval is required.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "City"
    table.rows[0].cells[1].text = "Limit (CNY/night)"
    for city, value in [("Beijing", "500"), ("Shanghai", "550")]:
        row = table.add_row()
        row.cells[0].text = city
        row.cells[1].text = value
    result = BytesIO()
    document.save(result)
    return result.getvalue()


def test_word_table_has_indexable_structural_evidence(document_factory):
    parsed, record = parsed_record(document_factory, word_bytes(), "travel.docx")
    assert parsed.tables[0].rows[1] == ["Shanghai", "550"]
    chunks = chunk_document(record, ChunkerConfig(mode="structure", chunk_size=500, overlap=0))
    assert any("Shanghai" in c.text and "CNY/night" in c.text for c in chunks if c.indexable)
    assert any("Approval is required" in c.text for c in chunks)


def test_html_table_has_indexable_structural_evidence(document_factory):
    body = (
        b"<h1>Freight</h1><p>Domestic only.</p><table><tr><th>Route</th><th>CNY/kg</th></tr>"
        b"<tr><td>North</td><td>12.50</td></tr></table>"
    )
    _, record = parsed_record(document_factory, body, "freight.html")
    chunks = chunk_document(record, ChunkerConfig(mode="structure", chunk_size=500, overlap=0))
    assert any("North" in c.text and "CNY/kg" in c.text for c in chunks)


def test_markdown_table_is_structured_without_reading_code_fence():
    body = (
        "# Policy\n\n| City | CNY/night |\n| --- | --- |\n| Shanghai | 550 |\n\n"
        "```text\n| Fake | value |\n| --- | --- |\n| ignored | 999 |\n```\n"
    )
    parsed = build_default_registry().parse_bytes(body.encode(), suffix=".md")
    assert len(parsed.tables) == 1
    assert parsed.tables[0].rows == [["Shanghai", "550"]]


def test_table_rows_repeat_headers_and_preserve_numbers(document_factory):
    body = "City,CNY/night\n" + "\n".join(f"City{i},{i}.50" for i in range(25))
    _, record = parsed_record(document_factory, body.encode(), "limits.csv")
    chunks = chunk_document(record, ChunkerConfig(mode="structure", chunk_size=100, overlap=0))
    tables = [c for c in chunks if c.kind == "table"]
    assert tables
    for i in range(25):
        assert sum(f"City{i} | {i}.50" in c.text.splitlines() for c in tables) == 1
    assert all("CNY/night" in c.text and len(c.text) <= 100 for c in tables)


def test_oversized_table_row_is_not_silently_truncated(document_factory):
    _, record = parsed_record(document_factory, b"Key,Value\nrow," + b"x" * 1000, "long.csv")
    with pytest.raises(ValueError, match="table_row_budget_exceeded"):
        chunk_document(record, ChunkerConfig(mode="structure", chunk_size=100, overlap=0))


def test_structural_chunks_keep_access_and_revision(document_factory):
    _, record = parsed_record(document_factory, b"Key,Amount\nrow,-1.50\n", "number.csv")
    chunks = chunk_document(record, ChunkerConfig(mode="structure", chunk_size=500, overlap=0))
    assert all(
        c.acl_groups == record.acl_groups and c.tenant_id == record.tenant_id for c in chunks
    )
    assert all(
        c.version_id == record.document_version.version_id and c.checksum == record.checksum
        for c in chunks
    )


def test_old_modes_have_not_been_silently_redefined(document_factory):
    _, record = parsed_record(document_factory, word_bytes(), "travel.docx")
    legacy = chunk_document(record, ChunkerConfig(mode="heading", chunk_size=500, overlap=0))
    assert not any("Shanghai" in c.text for c in legacy)
