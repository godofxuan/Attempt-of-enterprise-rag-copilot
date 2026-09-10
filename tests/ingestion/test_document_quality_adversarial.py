from io import BytesIO

import pytest
from docx import Document

from app.ingestion.chunking import ChunkerConfig, chunk_document
from tests.ingestion.test_document_quality_v1 import parsed_record


def test_docx_merged_cell_is_review_required_not_silently_flattened(document_factory):
    doc = Document()
    table = doc.add_table(rows=3, cols=2)
    for row, values in zip(
        table.rows, [["City", "Limit"], ["Shanghai", "550"], ["Beijing", "600"]], strict=True
    ):
        for cell, text in zip(row.cells, values, strict=True):
            cell.text = text
    table.cell(1, 0).merge(table.cell(2, 0))
    buffer = BytesIO()
    doc.save(buffer)
    _, record = parsed_record(document_factory, buffer.getvalue(), "merged.docx")
    assert any(w.code == "merged_table_cells" for w in record.parse_warnings)
    with pytest.raises(ValueError):
        chunk_document(record, ChunkerConfig(mode="structure"))


def test_html_merged_cell_is_review_required_even_with_equal_row_width(document_factory):
    html = (
        b'<table><tr><th>A</th><th>B</th></tr><tr><td rowspan="2">X</td><td>1</td></tr>'
        b"<tr><td>Y</td><td>2</td></tr></table>"
    )
    _, record = parsed_record(document_factory, html, "merged.html")
    assert any(w.code == "merged_table_cells" for w in record.parse_warnings)
    with pytest.raises(ValueError, match="review_required"):
        chunk_document(record, ChunkerConfig(mode="structure"))


def test_legacy_markdown_without_parsed_tables_keeps_raw_table(document_factory):
    text = "| Name | Value |\n| --- | --- |\n| No approval | -50 |"
    record = document_factory(text=text, parser_name="markdown", parser_version="1.0", tables=[])
    chunks = chunk_document(record, ChunkerConfig(mode="structure"))
    assert any("No approval" in c.text and "-50" in c.text for c in chunks)


def test_docx_row_locator_excludes_header(document_factory):
    from tests.ingestion.test_document_quality_v1 import word_bytes

    _, record = parsed_record(document_factory, word_bytes(), "travel.docx")
    chunks = chunk_document(record, ChunkerConfig(mode="structure", table_rows_per_chunk=1))
    table_chunks = [c for c in chunks if c.kind == "table"]
    assert table_chunks[0].locator.start == 2


def test_markdown_escaped_pipe_and_negative_sign_survive(document_factory):
    _, record = parsed_record(
        document_factory, b"| Name | CNY |\n| --- | --- |\n| A \\| B | -50.00 |\n", "table.md"
    )
    chunks = chunk_document(record, ChunkerConfig(mode="structure"))
    assert any("A | B" in c.text and "-50.00" in c.text for c in chunks)


def test_ragged_markdown_never_silently_drops_extra_cell(document_factory):
    with pytest.raises(Exception, match="row_width"):
        parsed_record(document_factory, b"| A | B |\n| --- | --- |\n| 1 | 2 | 3 |\n", "ragged.md")
